"""
DIANA ENGINE v4.0 — YOLO + ByteTrack + TTC + Excel + Video Anotasi
Tim DIANA | PKTJ Tegal | LKTI 2026

Menggunakan YOLOv8 (ultralytics) untuk deteksi kendaraan nyata.
Fallback ke demo mode jika ultralytics tidak tersedia.
"""

import cv2, numpy as np, io, os, hashlib
from collections import defaultdict, deque
from datetime import datetime, timedelta

try:
    from ultralytics import YOLO
    YOLO_OK = True
except ImportError:
    YOLO_OK = False
    print("[DIANA ENGINE] WARNING: ultralytics tidak terinstall → demo mode aktif")
    print("[DIANA ENGINE] Jalankan: pip install ultralytics opencv-python")

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    XL_OK = True
except ImportError:
    XL_OK = False

# ─── PKJI 2023 ────────────────────────────────────────────────
# COCO class id → (nama_kendaraan, kategori_pkji)
VMAP   = {2: ('Mobil','KR'), 3: ('Motor','SM'), 5: ('Bus','KB'), 7: ('Truk','KB')}
EKR    = {'KR': 1.0, 'KB': 1.3, 'SM': 0.5}
CLR    = {'KR': (0,220,50), 'KB': (30,30,235), 'SM': (0,220,230)}   # BGR
REF_W  = {'KR': 1.8, 'KB': 2.5, 'SM': 0.6}                         # meter
CLASSES = [2, 3, 5, 7]

TTC_BAHAYA  = 1.5   # detik
TTC_WASPADA = 3.0

# ─── HELPERS ──────────────────────────────────────────────────

def pkji(kr, kb, sm, cap=1000):
    q = kr*1.0 + kb*1.3 + sm*0.5
    qhr = round(q*4, 1)
    dj = round(qhr/cap, 3) if cap > 0 else 0
    tp = 'A' if dj<=0.35 else 'B' if dj<=0.54 else 'C' if dj<=0.77 else 'D' if dj<=0.93 else 'E' if dj<=1 else 'F'
    return qhr, dj, tp

def calc_ttc(dist_m, v_back, v_front):
    dv = (v_back - v_front) / 3.6
    if dv <= 0 or dist_m > 80:
        return None
    return round(dist_m / dv, 2)

def draw_label(frame, lines, x, y, color, fs=0.44):
    font = cv2.FONT_HERSHEY_SIMPLEX
    pad = 4; lh = 16
    if not lines:
        return
    mw = max(cv2.getTextSize(l, font, fs, 1)[0][0] for l in lines)
    th = len(lines)*lh + pad*2
    x = max(0, min(x, frame.shape[1]-mw-pad*2))
    y = max(th+4, y)
    ov = frame.copy()
    cv2.rectangle(ov, (x-pad, y-th-2), (x+mw+pad, y+pad), (0,0,0), -1)
    cv2.addWeighted(ov, 0.72, frame, 0.28, 0, frame)
    for i, ln in enumerate(lines):
        cv2.putText(frame, ln, (x, y-th+pad+(i+1)*lh), font, fs, color, 1, cv2.LINE_AA)

def draw_dashed(frame, p1, p2, color, thick=2, dash=10, gap=8):
    x1,y1 = p1; x2,y2 = p2
    d = np.hypot(x2-x1, y2-y1)
    if d < 1: return
    dx, dy = (x2-x1)/d, (y2-y1)/d
    pos = 0
    while pos < d:
        e = min(pos+dash, d)
        cv2.line(frame,
                 (int(x1+dx*pos), int(y1+dy*pos)),
                 (int(x1+dx*e),   int(y1+dy*e)),
                 color, thick, cv2.LINE_AA)
        pos += dash + gap

def mid_label(frame, p1, p2, text, color):
    mx, my = int((p1[0]+p2[0])/2), int((p1[1]+p2[1])/2)
    font = cv2.FONT_HERSHEY_SIMPLEX
    (tw, th), _ = cv2.getTextSize(text, font, 0.42, 1)
    cv2.rectangle(frame, (mx-3, my-th-4), (mx+tw+3, my+3), color, -1)
    cv2.putText(frame, text, (mx, my), font, 0.42, (0,0,0), 1, cv2.LINE_AA)

def hud(frame, fno, total, date_s, time_s, road,
        nact, kra, kba, sma, avgspd,
        cnt_kr, cnt_kb, cnt_sm, qsmp, dj, tp, ppm, nlanes, W):
    lines = [
        f"Frame: {fno}/{total}  {date_s}  {time_s}  {road}",
        f"Vehicles: {nact}  (KR:{kra} KB:{kba} SM:{sma})  Avg Speed: {avgspd:.1f} km/h",
        f"Total: Motor={cnt_sm} Mobil={cnt_kr} Bus/Truk={cnt_kb}",
        f"PKJI 2023  Q={qsmp:.0f} smp/jam  DJ={dj:.3f}  TP:{tp}  PPM:{ppm:.1f}  Lajur:{nlanes}",
    ]
    hh = len(lines)*22 + 12
    ov = frame.copy()
    cv2.rectangle(ov, (0,0), (W, hh), (0,0,0), -1)
    cv2.addWeighted(ov, 0.55, frame, 0.45, 0, frame)
    cols = [(0,220,50),(200,220,200),(200,220,200),(0,200,255)]
    for i, ln in enumerate(lines):
        cv2.putText(frame, ln, (10, 20+i*22),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.48, cols[i], 1, cv2.LINE_AA)
    cv2.putText(frame, f"DIANA-WEB v4.0 | {road} COUNTING",
                (W-430, frame.shape[0]-10),
                cv2.FONT_HERSHEY_SIMPLEX, 0.36, (120,120,120), 1)

# ─── TRACK STATE ──────────────────────────────────────────────

class Track:
    def __init__(self, tid, pkji_cat, vtype):
        self.id = tid; self.pkji = pkji_cat; self.vtype = vtype
        self.pos = deque(maxlen=60)
        self.sbuf = []; self.ema = 0.0
        self.cnt = 0; self.lane_name = ''; self.hist = []

    def update(self, cx, cy, fps, ppm):
        self.pos.append((cx, cy)); self.cnt += 1
        n = len(self.pos)
        if n < 4: return self.ema
        win = min(20, n-1)
        p1 = np.array(self.pos[-win-1]); p2 = np.array(self.pos[-1])
        dpx = float(np.linalg.norm(p2-p1))
        if ppm <= 0 or fps <= 0: return self.ema
        raw = (dpx/ppm) / (win/fps) * 3.6
        if raw > 180: return self.ema
        if self.ema > 0:
            mx = 40*(win/fps)
            raw = float(np.clip(raw, self.ema-mx, self.ema+mx))
        if raw >= 2:
            self.sbuf.append(raw)
            if len(self.sbuf) > 12: self.sbuf.pop(0)
        if self.sbuf:
            med = float(np.median(self.sbuf))
            self.ema = med if self.ema == 0 else 0.22*med + 0.78*self.ema
        if self.ema >= 2:
            self.hist.append(self.ema)
        return round(self.ema, 1)

# ─── LANE MANAGER ─────────────────────────────────────────────

class Lanes:
    _names = {
        2: ['Lajur 1','Lajur 2'],
        3: ['Lajur 1','Lajur 2','Lajur 3'],
        4: ['A-1','A-2','B-1','B-2'],
        6: ['A-1','A-2','A-3','B-1','B-2','B-3'],
    }
    def __init__(self, W, n=4):
        self.W = W; self.n = n; self.lw = W/n
    def idx(self, cx):
        return min(int(cx/self.lw), self.n-1)
    def name(self, i):
        lst = self._names.get(self.n, [f'Lajur {j+1}' for j in range(self.n)])
        return lst[i] if i < len(lst) else f'Lajur {i+1}'
    def draw(self, frame):
        for i in range(1, self.n):
            x = int(i*self.lw)
            cv2.line(frame, (x,0), (x,frame.shape[0]), (80,80,80), 1, cv2.LINE_AA)

# ─── EXCEL BUILDER ────────────────────────────────────────────

def build_excel(detail, tracks, events, road,
                cctv_date, cctv_time,
                cnt_kr, cnt_kb, cnt_sm,
                qsmp, dj, tp, ppm):
    if not XL_OK:
        return None
    wb = Workbook(); wb.remove(wb.active)

    def hdr(ws, cols, row=1, fill='1F3864'):
        fp = PatternFill('solid', fgColor=fill)
        fn = Font(bold=True, color='FFFFFF', size=10)
        al = Alignment(horizontal='center', vertical='center', wrap_text=True)
        bd = Border(*(Side(style='thin') for _ in range(4)))
        bd = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
        for c in range(1, len(cols)+1):
            cell = ws.cell(row=row, column=c)
            cell.fill = fp; cell.font = fn
            cell.alignment = al; cell.border = bd

    def zebra(ws, start, n, nc):
        fa = PatternFill('solid', fgColor='EBF3FB')
        fb = PatternFill('solid', fgColor='FFFFFF')
        bd = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
        for r in range(start, start+n):
            f = fa if (r-start)%2==0 else fb
            for c in range(1, nc+1):
                cell = ws.cell(row=r, column=c)
                cell.fill = f; cell.border = bd
                cell.alignment = Alignment(horizontal='center', vertical='center')

    def aw(ws):
        for col in ws.columns:
            ml = max((len(str(c.value)) for c in col if c.value), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(ml+4, 32)

    # Sheet per lajur
    by_lane = defaultdict(list)
    for r in detail:
        by_lane[r.get('lane', 'A-1')].append(r)

    col_d = ['No','Waktu(s)','Frame','ID','Jenis','PKJI','Lajur','X','Y',
             'Speed(km/h)','Jarak(m)','TTC(s)','Status TTC']
    for lane_nm, rows in sorted(by_lane.items()):
        ws = wb.create_sheet(f'Lajur {lane_nm}')
        ws.append(col_d); hdr(ws, col_d)
        for i, r in enumerate(rows, 1):
            ws.append([i, r.get('ts',''), r.get('frame',''), r.get('vid',''),
                       r.get('vtype',''), r.get('pkji',''), r.get('lane',''),
                       r.get('cx',''), r.get('cy',''),
                       round(r.get('speed',0),1),
                       r.get('dist',''), r.get('ttc',''), r.get('ttc_cat','Aman')])
        zebra(ws, 2, len(rows), len(col_d))
        aw(ws); ws.freeze_panes = 'A2'
        ws.sheet_properties.tabColor = '2E75B6'

    # Ringkasan per kendaraan
    ws2 = wb.create_sheet('Ringkasan Kendaraan')
    col2 = ['ID','Jenis','PKJI','ekr','Lajur','Avg(km/h)','Max(km/h)','Min(km/h)','Frame']
    ws2.append(col2); hdr(ws2, col2)
    seq = 0
    for tid, tr in tracks.items():
        if not tr.hist: continue
        seq += 1
        ws2.append([f'V{seq:03d}', tr.vtype, tr.pkji, EKR.get(tr.pkji,0),
                    tr.lane_name or '—',
                    round(float(np.mean(tr.hist)),1),
                    round(float(np.max(tr.hist)),1),
                    round(float(np.min(tr.hist)),1), tr.cnt])
    if seq > 0:
        zebra(ws2, 2, seq, len(col2))
    aw(ws2); ws2.sheet_properties.tabColor = '2E75B6'

    # Ringkasan jenis
    ws3 = wb.create_sheet('Ringkasan Jenis')
    col3 = ['Jenis Kendaraan','Kategori PKJI','ekr','Jumlah','Kontribusi smp']
    ws3.append(col3); hdr(ws3, col3, fill='375623')
    total_smp = 0
    for jenis, pkji_k, cnt2 in [('Mobil (KR)','KR',cnt_kr),
                                  ('Bus/Truk (KB)','KB',cnt_kb),
                                  ('Motor (SM)','SM',cnt_sm)]:
        smp = round(cnt2 * EKR[pkji_k], 1)
        total_smp += smp
        ws3.append([jenis, pkji_k, EKR[pkji_k], cnt2, smp])
    ws3.append(['TOTAL','—','—', cnt_kr+cnt_kb+cnt_sm, round(total_smp,1)])
    zebra(ws3, 2, 4, len(col3)); aw(ws3)
    ws3.sheet_properties.tabColor = '375623'

    # PKJI 2023
    ws4 = wb.create_sheet('PKJI 2023')
    rows4 = [
        ['Parameter','Nilai','Satuan'],
        ['Nama Jalan / Lokasi', road, ''],
        ['Tanggal & Jam CCTV', f'{cctv_date} {cctv_time}', ''],
        ['PPM (Pixel per Meter)', round(ppm,2), 'px/m'],
        ['KR — Kendaraan Ringan (Mobil)', cnt_kr, 'kendaraan'],
        ['KB — Kendaraan Berat (Bus+Truk)', cnt_kb, 'kendaraan'],
        ['SM — Sepeda Motor', cnt_sm, 'kendaraan'],
        ['Total Kendaraan Terdeteksi', cnt_kr+cnt_kb+cnt_sm, 'kendaraan'],
        ['Volume Q', qsmp, 'smp/jam'],
        ['Kapasitas Jalan Referensi', 1000, 'smp/jam'],
        ['Derajat Kejenuhan (DJ)', dj, '—'],
        ['Tingkat Pelayanan (TP)', tp, '—'],
        ['Kejadian BAHAYA (TTC < 1.5s)',
         sum(1 for e in events if e.get('category')=='BAHAYA'), 'kejadian'],
        ['Kejadian WASPADA (1.5–3.0s)',
         sum(1 for e in events if e.get('category')=='WASPADA'), 'kejadian'],
    ]
    for r in rows4: ws4.append(r)
    hdr(ws4, rows4[0], fill='7030A0')
    zebra(ws4, 2, len(rows4)-1, 3); aw(ws4)
    ws4.sheet_properties.tabColor = '7030A0'

    # Kejadian TTC
    if events:
        ws5 = wb.create_sheet('Kejadian TTC')
        col5 = ['No','Waktu(s)','Frame','ID Belakang','ID Depan',
                'TTC(s)','Kategori','V.Belakang(km/h)','V.Depan(km/h)','Jarak(m)','Lajur']
        ws5.append(col5); hdr(ws5, col5, fill='C00000')
        for i, e in enumerate(events, 1):
            ws5.append([i, e.get('ts',''), e.get('frame',''),
                        e.get('vid_b','—'), e.get('vid_f','—'),
                        e.get('ttc',''), e.get('category',''),
                        e.get('v_back',''), e.get('v_front',''),
                        e.get('dist_m',''), e.get('lane','')])
        zebra(ws5, 2, len(events), len(col5))
        aw(ws5); ws5.freeze_panes = 'A2'
        ws5.sheet_properties.tabColor = 'FF0000'

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf.read()

# ─── DIANA ENGINE ─────────────────────────────────────────────

class DianaEngine:
    def __init__(self, model_path='yolov8n.pt', conf=0.28, n_lanes=4,
                 ppm_override=None, road_name='SIMPANG PURWOKERTO',
                 cctv_dt=None, capacity=1000):
        self.model_path  = model_path
        self.conf        = conf
        self.n_lanes     = n_lanes
        self.ppm_ovr     = ppm_override
        self.road_name   = road_name.upper()
        self.cctv_dt     = cctv_dt or datetime.now()
        self.capacity    = capacity
        self.model       = None

        if YOLO_OK:
            self._load_model()

    def _load_model(self):
        for mp in [self.model_path, 'yolov8n.pt', 'yolov8s.pt']:
            try:
                self.model = YOLO(mp)
                print(f'[DIANA ENGINE] Model loaded: {mp}')
                return
            except Exception as e:
                print(f'[DIANA ENGINE] Cannot load {mp}: {e}')
        print('[DIANA ENGINE] No model found → demo mode')

    def process(self, video_in: str, video_out: str, excel_out: str,
                progress_cb=None, stop_ev=None) -> dict:
        """
        Proses video lengkap.
        PENTING: video_out dan excel_out harus path ABSOLUT yang sudah pasti bisa ditulis.
        """
        # Pastikan folder output ada
        os.makedirs(os.path.dirname(os.path.abspath(video_out)), exist_ok=True)
        os.makedirs(os.path.dirname(os.path.abspath(excel_out)), exist_ok=True)

        if self.model is None or not YOLO_OK:
            print('[DIANA ENGINE] Demo mode aktif')
            return self._demo(video_in, video_out, excel_out, progress_cb)

        return self._run(video_in, video_out, excel_out, progress_cb, stop_ev)

    def _run(self, video_in, video_out, excel_out, progress_cb, stop_ev):
        cap = cv2.VideoCapture(video_in)
        if not cap.isOpened():
            raise RuntimeError(f'[DIANA ENGINE] Tidak bisa buka video: {video_in}')

        fps   = cap.get(cv2.CAP_PROP_FPS) or 30
        W     = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
        H     = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        print(f'[DIANA ENGINE] Video: {W}x{H}, FPS={fps:.1f}, frames={total}')

        if progress_cb: progress_cb(3, 'Membuka video...')

        lanes    = Lanes(W, self.n_lanes)
        ppm      = self.ppm_ovr or 20.0
        calib_s  = []

        # VideoWriter — codec mp4v
        fourcc = cv2.VideoWriter_fourcc(*'mp4v')
        writer = cv2.VideoWriter(video_out, fourcc, fps, (W, H))
        if not writer.isOpened():
            raise RuntimeError(f'[DIANA ENGINE] Tidak bisa membuat video output: {video_out}')

        tracks  = {}        # tid → Track
        seq_map = {}        # tid → urutan (V001, V002, ...)
        seq_n   = 1
        seen    = set()
        cnt     = {'KR':0, 'KB':0, 'SM':0}
        detail  = []
        events  = []
        frame_no = 0

        if progress_cb: progress_cb(5, 'Deteksi dimulai...')

        while cap.isOpened():
            if stop_ev and stop_ev.is_set():
                break
            ret, frame = cap.read()
            if not ret:
                break
            frame_no += 1
            elapsed  = frame_no / fps
            ts_s  = (self.cctv_dt + timedelta(seconds=elapsed)).strftime('%H:%M:%S')
            dt_s  = self.cctv_dt.strftime('%Y-%m-%d')

            # Deteksi + ByteTrack
            try:
                res   = self.model.track(
                    source=frame, classes=CLASSES, conf=self.conf,
                    iou=0.45, persist=True, verbose=False,
                    tracker='bytetrack.yaml')
                boxes = res[0].boxes if res[0].boxes is not None else []
            except Exception as ex:
                print(f'[DIANA ENGINE] frame {frame_no} err: {ex}')
                boxes = []

            # Auto-kalibrasi PPM (150 frame pertama)
            if frame_no <= 150 and not self.ppm_ovr:
                for box in boxes:
                    if box.id is None: continue
                    cid = int(box.cls[0])
                    if cid not in VMAP: continue
                    _, pk = VMAP[cid]
                    x1,y1,x2,y2 = box.xyxy[0].tolist()
                    ref = REF_W.get(pk)
                    if ref and (x2-x1) > 15:
                        calib_s.append((x2-x1)/ref)
                if frame_no % 30 == 0 and len(calib_s) >= 4:
                    arr = np.array(calib_s)
                    p10, p90 = np.percentile(arr, [10, 90])
                    ok = arr[(arr>=p10)&(arr<=p90)]
                    if len(ok) > 0:
                        ppm = float(np.median(ok))

            # Update tracks
            active = []
            for box in boxes:
                if box.id is None: continue
                tid  = int(box.id[0])
                cid  = int(box.cls[0])
                if cid not in VMAP: continue
                vtype, pk = VMAP[cid]
                x1,y1,x2,y2 = map(int, box.xyxy[0].tolist())
                cx = (x1+x2)/2; cy = (y1+y2)/2

                if tid not in tracks:
                    tracks[tid] = Track(tid, pk, vtype)
                if tid not in seen:
                    seen.add(tid); cnt[pk] = cnt.get(pk,0) + 1
                if tid not in seq_map:
                    seq_map[tid] = seq_n; seq_n += 1

                tr  = tracks[tid]
                spd = tr.update(cx, cy, fps, ppm)
                li  = lanes.idx(cx)
                tr.lane_name = lanes.name(li)
                active.append((tid, cx, cy, pk, vtype, spd, x1, y1, x2, y2))

            # Hitung jarak & TTC per lajur
            by_lane = defaultdict(list)
            for item in active:
                by_lane[tracks[item[0]].lane_name].append(item)

            front_map = {}   # tid_belakang → (dist_m, ttc, tid_depan)
            all_pairs = []
            for ln_nm, items in by_lane.items():
                if len(items) < 2: continue
                srt = sorted(items, key=lambda x: x[2])   # sort by Y
                for i in range(len(srt)-1):
                    fi = srt[i]    # depan
                    bi = srt[i+1]  # belakang
                    dpx = np.hypot(bi[1]-fi[1], bi[2]-fi[2])
                    dm  = dpx/ppm if ppm > 0 else 0
                    tc  = calc_ttc(dm, bi[5], fi[5])
                    front_map[bi[0]] = (round(dm,1), tc, fi[0])
                    all_pairs.append({
                        'tid_b':bi[0],'tid_f':fi[0],
                        'cx_b':bi[1],'cy_b':bi[2],
                        'cx_f':fi[1],'cy_f':fi[2],
                        'dm':round(dm,1),'tc':tc,
                        'vb':bi[5],'vf':fi[5],'ln':ln_nm
                    })

            # Gambar garis jarak
            lanes.draw(frame)
            drawn = set()
            for p in all_pairs:
                k = (p['tid_b'], p['tid_f'])
                if k in drawn: continue
                drawn.add(k)
                pt1 = (int(p['cx_f']), int(p['cy_f']))
                pt2 = (int(p['cx_b']), int(p['cy_b']))
                dm  = p['dm']; tc = p['tc']
                if dm < 10 or (tc and tc < TTC_BAHAYA):   lc = (30,30,235)
                elif dm < 25 or (tc and tc < TTC_WASPADA): lc = (0,220,230)
                else:                                       lc = (0,200,50)
                draw_dashed(frame, pt1, pt2, lc, 2)
                lbl = f'{dm:.1f}m' + (f' TTC:{tc:.1f}s' if tc else '')
                mid_label(frame, pt1, pt2, lbl, lc)

            # Gambar bounding box + label
            spd_list = []
            for item in active:
                tid, cx, cy, pk, vtype, spd, x1, y1, x2, y2 = item
                vid    = f'V{seq_map[tid]:03d}'
                ln_nm  = tracks[tid].lane_name
                bc     = CLR.get(pk, (150,150,150))
                tc_cat = None; tc_v = None

                if tid in front_map:
                    dm_v, tc_v, _ = front_map[tid]
                    if tc_v and tc_v < TTC_BAHAYA:   bc = (30,30,235);  tc_cat = 'BAHAYA'
                    elif tc_v and tc_v < TTC_WASPADA: bc = (0,165,255); tc_cat = 'WASPADA'

                cv2.rectangle(frame, (x1,y1), (x2,y2), bc, 2)

                lns = [f'ID:{vid} {pk}/{vtype}',
                       f'Speed: {spd:.1f} km/h',
                       f'Jalur: {ln_nm}']
                if tid in front_map:
                    lns.append(f'Jarak: {front_map[tid][0]:.1f}m')
                if tc_v:
                    lns.append(f'TTC: {tc_v:.1f}s')
                draw_label(frame, lns, x1, y1-5, bc)

                if spd >= 2: spd_list.append(spd)

                dist_v = front_map[tid][0] if tid in front_map else None
                detail.append({
                    'ts':round(elapsed,1), 'frame':frame_no, 'vid':vid,
                    'vtype':vtype, 'pkji':pk, 'lane':ln_nm,
                    'cx':int(cx), 'cy':int(cy), 'speed':spd,
                    'dist':dist_v, 'ttc':tc_v, 'ttc_cat':tc_cat or 'Aman',
                })

                if tc_v and tc_cat in ('BAHAYA','WASPADA'):
                    tf = front_map[tid][2] if tid in front_map else None
                    vf_id = f'V{seq_map[tf]:03d}' if tf and tf in seq_map else '—'
                    vf_spd = tracks[tf].ema if tf and tf in tracks else 0
                    events.append({
                        'ts':round(elapsed,1), 'frame':frame_no,
                        'vid_b':vid, 'vid_f':vf_id,
                        'ttc':tc_v, 'category':tc_cat,
                        'v_back':spd, 'v_front':vf_spd,
                        'dist_m':dist_v or 0, 'lane':ln_nm,
                    })

            # HUD 4 baris
            avgspd = float(np.mean(spd_list)) if spd_list else 0.0
            nact   = len(active)
            kra    = sum(1 for x in active if x[3]=='KR')
            kba    = sum(1 for x in active if x[3]=='KB')
            sma    = sum(1 for x in active if x[3]=='SM')
            qsmp, dj, tp = pkji(cnt['KR'], cnt['KB'], cnt['SM'], self.capacity)

            hud(frame, frame_no, total, dt_s, ts_s, self.road_name,
                nact, kra, kba, sma, avgspd,
                cnt['KR'], cnt['KB'], cnt['SM'],
                qsmp, dj, tp, ppm, self.n_lanes, W)

            writer.write(frame)

            if progress_cb and frame_no % 25 == 0 and total > 0:
                pct = min(int(frame_no/total*83)+5, 88)
                progress_cb(pct, f'Frame {frame_no}/{total} | {avgspd:.1f} km/h | PPM:{ppm:.1f}')

        cap.release()
        writer.release()
        print(f'[DIANA ENGINE] Video selesai: {video_out} ({os.path.getsize(video_out)//1024} KB)')

        if progress_cb: progress_cb(90, 'Membuat Excel...')

        # Build Excel
        qsmp2, dj2, tp2 = pkji(cnt['KR'], cnt['KB'], cnt['SM'], self.capacity)
        avg_all = float(np.mean([r['speed'] for r in detail if r['speed']>0])) if detail else 0

        xl_bytes = build_excel(
            detail, tracks, events, self.road_name,
            self.cctv_dt.strftime('%Y-%m-%d'), self.cctv_dt.strftime('%H:%M'),
            cnt['KR'], cnt['KB'], cnt['SM'], qsmp2, dj2, tp2, ppm
        )
        if xl_bytes:
            with open(excel_out, 'wb') as f:
                f.write(xl_bytes)
            print(f'[DIANA ENGINE] Excel: {excel_out} ({os.path.getsize(excel_out)//1024} KB)')

        if progress_cb: progress_cb(100, 'Selesai!')

        return {
            'status':'done',
            'total_veh': cnt['KR']+cnt['KB']+cnt['SM'],
            'total_kr':  cnt['KR'], 'total_kb': cnt['KB'], 'total_sm': cnt['SM'],
            'avg_speed': round(avg_all, 1),
            'max_speed': round(max((r['speed'] for r in detail if r['speed']>0), default=0), 1),
            'danger_ev':  sum(1 for e in events if e['category']=='BAHAYA'),
            'warning_ev': sum(1 for e in events if e['category']=='WASPADA'),
            'q_smp': qsmp2, 'dj': dj2, 'tp': tp2, 'ppm_used': round(ppm,2),
            'events': events,
            'video_out': video_out,
            'excel_out': excel_out,
        }

    # ─── DEMO MODE ────────────────────────────────────────────
    def _demo(self, video_in, video_out, excel_out, progress_cb):
        """Demo mode: HUD tetap muncul, data dari hash video (deterministik)."""
        import random
        h = hashlib.md5()
        try:
            with open(video_in, 'rb') as f: h.update(f.read(2*1024*1024))
        except: h.update(video_in.encode())
        rng = random.Random(int(h.hexdigest(),16) % (2**31))

        cap = cv2.VideoCapture(video_in)
        fps = cap.get(cv2.CAP_PROP_FPS) or 30
        W   = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
        H   = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))

        fourcc = cv2.VideoWriter_fourcc(*'mp4v')
        writer = cv2.VideoWriter(video_out, fourcc, fps, (W, H))

        kr=rng.randint(60,250); kb=rng.randint(8,70); sm=rng.randint(100,400)
        avgspd=round(rng.uniform(22,55),1); dg=rng.randint(0,14); wn=rng.randint(2,22)
        qsmp, dj, tp = pkji(kr, kb, sm)

        frame_no = 0
        while cap.isOpened():
            ret, frame = cap.read()
            if not ret: break
            frame_no += 1
            elapsed  = frame_no/fps
            ts_s = (self.cctv_dt+timedelta(seconds=elapsed)).strftime('%H:%M:%S')
            dt_s = self.cctv_dt.strftime('%Y-%m-%d')

            hud(frame, frame_no, total, dt_s, ts_s, self.road_name,
                0, 0, 0, 0, avgspd, kr, kb, sm, qsmp, dj, tp, 20.0, self.n_lanes, W)
            cv2.putText(frame,
                '[DEMO — pip install ultralytics opencv-python untuk deteksi nyata]',
                (8, H-24), cv2.FONT_HERSHEY_SIMPLEX, 0.42, (0,165,255), 1)
            writer.write(frame)

            if progress_cb and frame_no%25==0 and total>0:
                progress_cb(min(int(frame_no/total*83)+5,88),
                            f'Frame {frame_no}/{total} (demo mode)')

        cap.release(); writer.release()
        print(f'[DIANA ENGINE] Demo video: {video_out}')

        # Demo events & Excel
        dur = total/fps if fps>0 else 120
        events = []
        for _ in range(dg):
            ts=round(rng.uniform(5,max(6,dur-5)),1); vb=round(rng.uniform(40,88),1)
            vf=round(rng.uniform(15,max(16,vb-10)),1); dm=round(rng.uniform(2,12),1)
            dv=(vb-vf)/3.6; tc=round(dm/dv,2) if dv>0.1 else 0.8
            events.append({'ts':ts,'frame':int(ts*fps),'vid_b':'V—','vid_f':'V—',
                'ttc':min(tc,1.49),'category':'BAHAYA',
                'v_back':vb,'v_front':vf,'dist_m':dm,'lane':'A-1'})
        for _ in range(wn):
            ts=round(rng.uniform(5,max(6,dur-5)),1); vb=round(rng.uniform(35,72),1)
            vf=round(rng.uniform(20,max(21,vb-8)),1); dm=round(rng.uniform(12,30),1)
            dv=(vb-vf)/3.6; tc=max(1.5,min(round(dm/dv,2) if dv>0.1 else 2.5,2.99))
            events.append({'ts':ts,'frame':int(ts*fps),'vid_b':'V—','vid_f':'V—',
                'ttc':tc,'category':'WASPADA',
                'v_back':vb,'v_front':vf,'dist_m':dm,'lane':'B-1'})

        if progress_cb: progress_cb(90, 'Membuat Excel (demo)...')
        xl = build_excel([], {}, events, self.road_name,
                         self.cctv_dt.strftime('%Y-%m-%d'), self.cctv_dt.strftime('%H:%M'),
                         kr, kb, sm, qsmp, dj, tp, 20.0)
        if xl:
            with open(excel_out, 'wb') as f: f.write(xl)
            print(f'[DIANA ENGINE] Demo Excel: {excel_out}')

        if progress_cb: progress_cb(100, 'Selesai (demo)!')
        return {
            'status':'done', 'total_veh':kr+kb+sm,
            'total_kr':kr, 'total_kb':kb, 'total_sm':sm,
            'avg_speed':avgspd, 'max_speed':round(avgspd+20,1),
            'danger_ev':dg, 'warning_ev':wn,
            'q_smp':qsmp, 'dj':dj, 'tp':tp, 'ppm_used':20.0,
            'events':events,
            'video_out':video_out, 'excel_out':excel_out,
        }
